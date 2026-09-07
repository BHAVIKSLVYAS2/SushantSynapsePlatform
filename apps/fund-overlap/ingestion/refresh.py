"""Build audited, immutable PPFAS equity snapshots; publish index only on full success."""
import argparse
import datetime as dt
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import tempfile
import urllib.parse
import urllib.request
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INDEX_URL = 'https://amc.ppfas.com/downloads/portfolio-disclosure/index.php'
FUNDS = {
    'PPFCF': ('ppfas-flexi-cap', 'Parag Parikh Flexi Cap Fund', 'Flexi Cap'),
    'PPTSF': ('ppfas-elss', 'Parag Parikh ELSS Tax Saver Fund', 'ELSS'),
    'PPLCF': ('ppfas-large-cap', 'Parag Parikh Large Cap Fund', 'Large Cap'),
    'PPCHF': ('ppfas-conservative-hybrid', 'Parag Parikh Conservative Hybrid Fund', 'Conservative Hybrid'),
    'PPDAAF': ('ppfas-dynamic-allocation', 'Parag Parikh Dynamic Asset Allocation Fund', 'Dynamic Asset Allocation'),
}

def valid_isin(value):
    if not re.fullmatch(r'[A-Z]{2}[A-Z0-9]{9}[0-9]', value):
        return False
    digits = ''.join(str(ord(c) - 55) if c.isalpha() else c for c in value)
    total = 0
    for i, digit in enumerate(reversed(digits)):
        n = int(digit) * (2 if i % 2 else 1)
        total += n - 9 if n > 9 else n
    return total % 10 == 0

def discover(html):
    result = {}
    for href in re.findall(r'href=[\"\']([^\"\']+\.xlsx[^\"\']*)', html):
        url = urllib.parse.urljoin(INDEX_URL, href.replace('&amp;', '&'))
        if urllib.parse.urlparse(url).hostname != 'amc.ppfas.com':
            continue
        match = re.search(r'/([A-Z]+)_PPFAS_Monthly_Portfolio_Report_([A-Za-z]+_\d{1,2}_\d{4})\.xlsx', url)
        if not match or match[1] not in FUNDS:
            continue
        date = dt.datetime.strptime(match[2], '%B_%d_%Y').date()
        if date > dt.datetime.now(dt.timezone.utc).date():
            continue
        if match[1] not in result or date > result[match[1]][0]:
            result[match[1]] = (date, url)
    if set(result) != set(FUNDS):
        raise ValueError('Official index is missing supported funds; previous index retained')
    return result

def parse_ppfas(raw, code, expected_date, url, fetched_at):
    workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    if code not in workbook.sheetnames:
        raise ValueError('Unexpected workbook sheet')
    rows = list(workbook[code].iter_rows())
    heading = ' '.join(str(c.value or '') for row in rows[:4] for c in row)
    match = re.search(r'as on\s+([A-Za-z]+\s+\d{1,2},\s*\d{4})', heading, re.I)
    if not match or dt.datetime.strptime(re.sub(r'\s+', ' ', match[1]), '%B %d, %Y').date() != expected_date:
        raise ValueError('Workbook portfolio date does not match its URL')
    if FUNDS[code][1].lower() not in heading.lower():
        raise ValueError('Workbook fund name mismatch')
    header = next((row for row in rows[:12] if any(str(c.value).strip() == 'ISIN' for c in row)), None)
    if not header or str(header[2].value).strip() != 'ISIN' or 'Net' not in str(header[6].value):
        raise ValueError('Disclosure columns changed')
    included, section, grand_total, below_precision = {}, False, None, 0
    for row in rows[4:]:
        name = str(row[1].value or '').strip()
        isin = re.sub(r'\s+', '', str(row[2].value or '')).upper()
        cell = row[6]
        low = name.lower()
        if 'grand total' in low:
            grand_total = cell.value
            break
        if low.startswith('equity & equity related'):
            section = True
        elif any(term in low for term in ['reits', 'invits', 'unlisted', 'money market instruments', 'debt instruments', 'debt &', 'securitised debt', 'mutual fund units', 'others', 'derivatives']):
            section = False
        if section and isin and isinstance(cell.value, (int, float)) and cell.value > 0 and not re.fullmatch(r'[A-Z]{2}[A-Z0-9]{9}[0-9]', isin):
            raise ValueError('Malformed ISIN on a weighted equity row')
        if not section or not re.fullmatch(r'[A-Z]{2}[A-Z0-9]{9}[0-9]', isin):
            continue
        if not valid_isin(isin):
            raise ValueError('Invalid ISIN check digit: ' + isin)
        value = cell.value
        if isinstance(value, str) and value.strip() == '$0.00%':
            below_precision += 1
            continue
        if not isinstance(value, (int, float)) or isinstance(value, bool) or not math.isfinite(value) or '%' not in cell.number_format:
            raise ValueError('Missing numeric percentage for ' + isin)
        weight = round(value * 100, 6)
        if weight < 0 or weight > 100:
            raise ValueError('Invalid equity weight')
        if not weight:
            continue
        sector = str(row[3].value or 'Unknown').strip()
        if isin in included:
            if included[isin]['sector'] != sector:
                raise ValueError('Conflicting industry for ' + isin)
            included[isin]['weight'] = round(included[isin]['weight'] + weight, 6)
        else:
            included[isin] = dict(isin=isin, name=name, sector=sector, weight=weight)
    workbook.close()
    total = round(sum(h['weight'] for h in included.values()), 6)
    if not isinstance(grand_total, (int, float)) or abs(grand_total - 1) > .005:
        raise ValueError('Grand total is not 100%')
    if not 0 < total <= 100.5 or len(included) < 5:
        raise ValueError('Invalid included equity coverage')
    scheme, name, category = FUNDS[code]
    return dict(schemeId=scheme, name=name, amc='PPFAS Mutual Fund', category=category,
                portfolioDate=expected_date.isoformat(), includedNavWeight=total, belowPrecisionCount=below_precision,
                policy='Physical listed equities including foreign and arbitrage legs; excludes debt, cash, derivatives, fund units and REIT/InvIT units. No hedge netting.',
                source=dict(publisher='PPFAS Mutual Fund', url=url, indexUrl=INDEX_URL,
                            fetchedAt=fetched_at, sha256=hashlib.sha256(raw).hexdigest(), adapterVersion=1),
                holdings=sorted(included.values(), key=lambda h: (-h['weight'], h['isin'])))

def download(url):
    request = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0', 'Accept': '*/*'})
    with urllib.request.urlopen(request, timeout=60) as response:
        raw = response.read(20_000_001)
    if len(raw) > 20_000_000:
        raise ValueError('Source file too large')
    return raw

def atomic_write(path, raw):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(dir=path.parent, suffix='.tmp')
    try:
        with os.fdopen(fd, 'wb') as handle:
            handle.write(raw)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)

def refresh(output, source_dir=None, index_file=None):
    html = Path(index_file).read_text(encoding='utf-8-sig') if index_file else download(INDEX_URL).decode()
    sources = discover(html)
    now = dt.datetime.now(dt.timezone.utc).isoformat()
    previous_path = output / 'fund-index.json'
    previous = json.loads(previous_path.read_text()) if previous_path.exists() else {'funds': []}
    old = {f['id']: f for f in previous['funds']}
    pending, entries = [], []
    for code, (date, url) in sources.items():
        raw = (Path(source_dir) / (code + '.xlsx')).read_bytes() if source_dir else download(url)
        snapshot = parse_ppfas(raw, code, date, url, now)
        prior = old.get(snapshot['schemeId'])
        if prior and date.isoformat() < prior['portfolioDate']:
            raise ValueError('Refusing to regress the published date')
        if prior and prior['sourceSha256'] == snapshot['source']['sha256']:
            entries.append(prior)
            continue
        # Include content hash: corrected disclosures create a new immutable version.
        encoded = (json.dumps(snapshot, indent=2, ensure_ascii=False) + '\n').encode()
        digest = hashlib.sha256(encoded).hexdigest()
        relative = 'holdings/' + snapshot['schemeId'] + '/' + date.isoformat() + '-' + digest[:16] + '.json'
        path = output / relative
        if path.exists() and path.read_bytes() != encoded:
            raise ValueError('Immutable snapshot collision')
        pending.append((path, encoded))
        entries.append(dict(id=snapshot['schemeId'], name=snapshot['name'], amc=snapshot['amc'], category=snapshot['category'], portfolioDate=snapshot['portfolioDate'], includedNavWeight=snapshot['includedNavWeight'], holdingsCount=len(snapshot['holdings']), holdingsPath='/api/fund-overlap/' + relative, sourceSha256=snapshot['source']['sha256'], snapshotSha256=digest))
    entries.sort(key=lambda f: f['name'])
    if entries == previous['funds']:
        print('No disclosure changes; last valid index retained')
        return
    index = dict(version=1, generatedAt=now, coverage='Five PPFAS equity/hybrid schemes. Other AMCs are not yet supported.', funds=entries)
    # Nothing is published before every source has passed validation.
    for path, raw in pending:
        atomic_write(path, raw)
    atomic_write(previous_path, (json.dumps(index, indent=2) + '\n').encode())
    for fund in entries:
        print(f"{fund['name']}: {fund['portfolioDate']}, {fund['holdingsCount']} equities, {fund['includedNavWeight']}% NAV")

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, default=ROOT / 'data')
    parser.add_argument('--source-dir', type=Path)
    parser.add_argument('--index-file', type=Path)
    args = parser.parse_args()
    refresh(args.output, args.source_dir, args.index_file)
