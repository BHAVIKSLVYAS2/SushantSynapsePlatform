"""Conservative official monthly portfolio reader. No identifier or weight guessing.

Each worksheet must identify one scheme, its reporting month, an ISIN column,
an explicitly labelled % NAV column and instrument sections. Unsupported or
ambiguous worksheets are reported; they never become an equity snapshot.
"""
import calendar
import datetime as dt
import hashlib
import io
import math
import re
import zipfile
from dataclasses import dataclass

from openpyxl import load_workbook


POLICY = ('Physical listed equities including foreign and arbitrage legs; excludes debt, '
          'cash, derivatives, fund units and REIT/InvIT units. No hedge netting.')
MAX_DOCUMENT = 40_000_000
MAX_EXPANDED = 100_000_000
MAX_ZIP_MEMBERS = 200


@dataclass(frozen=True)
class Cell:
    value: object
    number_format: str = 'General'


def text(value):
    return re.sub(r'\s+', ' ', str(value if value is not None else '')).strip()


def valid_isin(value):
    if not re.fullmatch(r'[A-Z]{2}[A-Z0-9]{9}\d', value):
        return False
    digits = ''.join(str(ord(c) - 55) if c.isalpha() else c for c in value)
    return sum((lambda n: n - 9 if n > 9 else n)(int(v) * (2 if i % 2 else 1))
               for i, v in enumerate(reversed(digits))) % 10 == 0


def family_name(name):
    """Remove only recognised trailing plan/options, never interior fund words.

    Direct and regular plans hold the same portfolio. We deliberately do not
    fuzzy-match renamed funds, AMC abbreviations, or similar investment styles.
    """
    value = text(name).lower().replace('–', '-').replace('—', '-').replace('&', ' and ')
    value = re.sub(r'\s*\((?:formerly|erstwhile).*$', '', value)
    # The first explicitly delimited plan or distribution option begins the
    # suffix. A fund named Growth Opportunities keeps those interior words.
    value = re.split(r'\s*-\s*(?=(?:direct|regular|retail|institutional)\b(?:\s+plan)?'
                     r'(?:\s*(?:-|$)|\s+(?:growth|idcw|dividend)\b)|'
                     r'(?:growth|idcw|dividend)\b(?=\s*(?:-|$|\(|option\b|payout\b|reinvestment\b)))', value)[0]
    value = re.sub(r'\s+(?:(?:direct|regular|retail|institutional)(?:\s+plan)?\s*)'
                   r'(?:(?:growth|idcw|dividend|option|payout|reinvestment)[\s()-]*)*$', '', value)
    value = re.sub(r'\s+(?:growth|idcw|dividend)(?:\s+(?:option|payout|reinvestment|plan))*$', '', value)
    return re.sub(r'[^a-z0-9]+', ' ', value).strip()


def catalogue_families(catalogue):
    families = {}
    for item in catalogue:
        if not isinstance(item, dict) or not isinstance(item.get('schemeName'), str):
            raise ValueError('Invalid fund catalogue item')
        code = item.get('schemeCode')
        if isinstance(code, bool) or not str(code).isdigit():
            raise ValueError('Invalid catalogue scheme code')
        key = family_name(item['schemeName'])
        families.setdefault(key, []).append(dict(schemeCode=int(code), name=item['schemeName']))
    return families


def source_prefixes(source):
    prefixes = source.get('schemeNamePrefixes') or source.get('namePrefixes')
    if prefixes:
        return [family_name(p) for p in prefixes]
    amc = re.sub(r'\s+(?:Mutual Fund|Asset Management.*)$', '', source['amc'], flags=re.I)
    if source.get('id') == 'ppfas':
        return ['parag parikh', 'ppfas']
    return [family_name(amc)]


def owns_family(source, family):
    return any(family == p or family.startswith(p + ' ') for p in source_prefixes(source))


def parse_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    value = text(value)
    value = re.sub(r',\s*', ', ', value)
    value = re.sub(r'(\d)(?:st|nd|rd|th)\b', r'\1', value, flags=re.I)
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%d.%m.%Y',
                '%B %d, %Y', '%b %d, %Y', '%d %B %Y', '%d %b %Y',
                '%d-%B-%Y', '%d-%b-%Y', '%B %d %Y', '%b %d %Y', '%d.%B.%Y', '%d.%b.%Y'):
        try:
            return dt.datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ValueError('Unrecognised portfolio date')


def heading_date(rows, expected_date):
    candidates = set()
    joined = ' '.join(text(c.value) for row in rows[:30] for c in row)
    for match in re.finditer(r'(?:as\s+(?:on|at|of)|portfolio\s+(?:date|for)|(?:month|period)\s+ended)\s*[:\-]?\s*'
                             r'(\d{4}-\d{2}-\d{2}|\d{1,2}[-/.]\d{1,2}[-/.]\d{4}|'
                             r'\d{1,2}(?:st|nd|rd|th)?[\s.-]+[A-Za-z]+[\s,.-]+\d{4}|'
                             r'[A-Za-z]+\s+\d{1,2}(?:st|nd|rd|th)?,?\s*\d{4})', joined, re.I):
        candidates.add(parse_date(match[1]))
    for row in rows[:15]:
        for cell in row:
            value = text(cell.value)
            if re.search(r'monthly\s+portfolio', value, re.I):
                match = re.search(r'\bfor\s+([A-Za-z]+)\s+(\d{4})\b', value, re.I)
                if match:
                    first = dt.datetime.strptime(match[1] + ' ' + match[2], '%B %Y').date()
                    candidates.add(first.replace(day=calendar.monthrange(first.year, first.month)[1]))
    # A labelled date cell may use an Excel date rather than rendered text.
    for row in rows[:20]:
        if any(re.search(r'portfolio date|as on|as at|as of|month ended', text(c.value), re.I) for c in row):
            candidates.update(c.value.date() if isinstance(c.value, dt.datetime) else c.value
                              for c in row if isinstance(c.value, (dt.datetime, dt.date)))
    if len(candidates) != 1:
        raise ValueError('A unique labelled portfolio date is required')
    date = candidates.pop()
    if date != expected_date:
        raise ValueError('Workbook portfolio date does not match the discovered month')
    if date.day != calendar.monthrange(date.year, date.month)[1]:
        raise ValueError('Monthly portfolio date must be a calendar month end')
    return date


def _candidate_name(value):
    value = text(value)
    value = re.sub(r'^(?:name\s+of\s+(?:the\s+)?scheme|scheme\s+name)\s*[:\-]\s*', '', value, flags=re.I)
    value = re.sub(r'^monthly\s+portfolio\s+(?:statement\s+)?of\s+', '', value, flags=re.I)
    value = re.sub(r'^(?=[A-Z0-9]*\d)[A-Z0-9]{2,10}\s*-\s*', '', value)
    value = re.sub(r'\s*[-:]?\s*an?\s+open[- ]ended.*$', '', value, flags=re.I)
    value = re.sub(r'\s*[-:]?\s*(?:monthly\s+portfolio|portfolio\s+(?:of|as|for)|as\s+(?:on|at|of)).*$', '', value, flags=re.I)
    value = re.sub(r'\s*\((?:an?\s+open|an?\s+close|formerly|erstwhile).*$', '', value, flags=re.I)
    value = re.sub(r'\s+for\s+[A-Za-z]+\s+\d{4}.*$', '', value, flags=re.I)
    return value.strip(' :-(')


def scheme_identity(rows, source, families):
    candidates = {}
    aliases = {family_name(k): v for k, v in source.get('schemeAliases', {}).items()}
    for row in rows[:20]:
        for cell in row:
            for fragment in str(cell.value or '').split('\n'):
                name = _candidate_name(fragment)
                key = family_name(name)
                if not name or len(name) > 180 or not re.search(r'\bfund\b|\betf\b', key):
                    continue
                key = family_name(aliases.get(key, name))
                if owns_family(source, key) and key != family_name(source['amc']):
                    candidates[key] = aliases.get(key, name)
    exact = {key: name for key, name in candidates.items() if key in families}
    chosen = exact or candidates
    if len(chosen) != 1:
        raise ValueError('A unique official scheme heading and AMC identity are required')
    key, name = next(iter(chosen.items()))
    if key not in families:
        # Spacing/punctuation changes (AngelOne/Angel One, Nifty50/Nifty 50)
        # are accepted only if the complete preserved name is unambiguous.
        exact_compact = [f for f in families if f.replace(' ', '') == key.replace(' ', '')]
        if len(exact_compact) == 1:
            key = exact_compact[0]
    source_slug = source['id'][:36]
    slug = re.sub(r'[^a-z0-9]+', '-', key)[:61 - len(source_slug)].strip('-')
    identity_hash = hashlib.sha256((source['id'] + ':' + key).encode()).hexdigest()[:8]
    return dict(id=f"official-{source_slug}-{slug}-{identity_hash}", name=name,
                family=key, schemeCodes=sorted({v['schemeCode'] for v in families.get(key, [])}))


def unsupported_reason(name):
    # An equity ETF is supported; debt/gold ETFs and fund-of-funds are not.
    low = family_name(name)
    if re.search(r'\b(?:fund of funds|fund of fund|fof)\b', low):
        return 'Fund-of-funds portfolios are not comparable as direct equity holdings.'
    if re.search(r'\b(?:gold|silver|liquid|overnight|money market|gilt|bond|debt|credit risk|'
                 r'ultra short|low duration|short duration|medium duration|long duration|corporate bond|'
                 r'banking and psu|fixed maturity|fmp)\b', low):
        return 'Debt, cash and precious-metal portfolios are outside physical-equity comparison.'
    return None


def section_type(name):
    low = text(name).lower().replace('&', ' and ')
    low = re.sub(r'^\s*(?:\([a-z0-9]{1,3}\)|[a-z0-9]{1,3}[.)])\s*', '', low)
    # Unlisted securities and derivative legs must never enter physical equities.
    if re.search(r'\b(unlisted|derivatives?|futures?|options?|reits?|invits?|mutual fund|fund units|'
                 r'exchange traded fund|etf|debt|bonds?|debentures?|money market|cash|treasury|'
                 r'commercial paper|certificate|securiti[sz]ed|gold|silver|repo|treps|others?|'
                 r'net (?:current|receivables)|receivables|payables|preference shares|real est inv trust|infra inv trust|deposits)\b', low):
        return 'excluded'
    if re.search(r'\bequit(?:y|ies)\b|\b(?:domestic|foreign|overseas|international) shares\b', low):
        return 'equity'
    if re.search(r'listed|awaiting listing', low):
        return 'listed'
    return None


def section_transition(kind, label, parent):
    low = label.lower()
    if kind == 'listed':
        if parent is None:
            raise ValueError('Listing status has no instrument classification')
        return parent, parent
    if re.search(r'\bunlisted\b', low) and not re.search(r'\b(?:debt|bonds?|debentures?|money market|fund units|mutual fund|reits?|invits?)\b', low):
        # Listing subsections change inclusion, while preserving asset class.
        return 'excluded', 'equity' if re.search(r'\bequit(?:y|ies)\b', low) else parent
    return kind, kind


def find_columns(rows):
    for index, row in enumerate(rows[:35]):
        isin_columns = [i for i, c in enumerate(row) if re.fullmatch(r'isin(?:\s+code)?', text(c.value), re.I)]
        if not isin_columns:
            continue
        if len(isin_columns) != 1:
            raise ValueError('Ambiguous ISIN columns')
        labels = []
        for col in range(len(row)):
            labels.append(' '.join(text(rows[r][col].value) for r in range(max(0, index - 1), min(len(rows), index + 3))
                                   if col < len(rows[r])).lower())
        def is_weight(label):
            return '%' in label and re.search(r'\bnav\b|\baum\b|net\s+assets|net\s+asset\s+value', label) and not re.search(r'derivative|hedged|benchmark', label)
        weights = [i for i, c in enumerate(row) if is_weight(text(c.value).lower())]
        if not weights:
            weights = [i for i, label in enumerate(labels) if is_weight(label)]
        if len(weights) != 1:
            raise ValueError('One explicitly labelled % of NAV column is required')
        names = [i for i, c in enumerate(row) if re.search(r'^(?:name|particulars|instrument|security|company/issuer/instrument name)\b', text(c.value), re.I)
                 and not re.search(r'industry|sector|rating', text(c.value), re.I)]
        sectors = [i for i, c in enumerate(row) if re.search(r'industry|sector', text(c.value), re.I)]
        if len(names) != 1:
            raise ValueError('A unique security-name column is required')
        return index, isin_columns[0], weights[0], names[0], sectors[0] if len(sectors) == 1 else None
    raise ValueError('No recognised portfolio column header')


def weight_value(cell):
    value = cell.value
    if isinstance(value, str):
        stripped = value.strip()
        if re.fullmatch(r'[<>]?\s*0(?:\.0+)?\s*%', stripped):
            return 0.0, 'percent-string'
        match = re.fullmatch(r'(-?\d+(?:\.\d+)?)\s*%', stripped)
        if not match:
            raise ValueError('Non-numeric percentage in a weighted row')
        result, units = float(match[1]), 'percent-string'
    elif isinstance(value, (float, int)) and not isinstance(value, bool) and math.isfinite(value):
        units = 'excel-percent' if '%' in cell.number_format else 'nav-percent'
        result = value * 100 if units == 'excel-percent' else value
    else:
        raise ValueError('Missing or non-finite NAV percentage')
    if not -100.5 <= result <= 100.5:
        raise ValueError('NAV percentage outside valid bounds')
    return round(result, 6), units


def parse_sheet(name, rows, source, document, families, fetched_at, source_sha):
    header, isin_col, weight_col, name_col, sector_col = find_columns(rows)
    identity = scheme_identity(rows[:header], source, families)
    date = heading_date(rows[:header + 1], parse_date(document['portfolioDate']))
    reason = unsupported_reason(identity['name'])
    if reason:
        return None, dict(**identity, portfolioDate=date.isoformat(), message=reason)
    included, unit_types, grand_totals, equity_totals = {}, set(), [], []
    section, parent, seen_equity, weighted_equity, omitted_zero = None, None, False, False, 0
    for row in rows[header + 1:]:
        if len(row) <= max(name_col, isin_col, weight_col):
            continue
        security = text(row[name_col].value)
        isin = re.sub(r'\s+', '', text(row[isin_col].value)).upper()
        raw_weight = row[weight_col].value
        # HDFC places section labels in the ISIN column, security names in a
        # later column. Only recognise explicit classifications in that slot.
        if not security and isin and not valid_isin(isin):
            classification = text(row[isin_col].value)
            if section_type(classification) or re.fullmatch(r'(?:sub\s*)?total|grand\s+total.*', classification, re.I):
                security, isin = classification, ''
        low = security.lower()
        if re.fullmatch(r'isin(?:\s+code)?', isin, re.I):
            raise ValueError('Multiple schemes/headers in one worksheet need a dedicated adapter')
        if re.search(r'grand\s+total|total\s+(?:net\s+assets|portfolio)', low) and not isin:
            value, units = weight_value(row[weight_col]); unit_types.add(units)
            grand_totals.append(value)
            break
        if re.fullmatch(r'(?:sub\s*)?total', low) and not isin:
            if re.fullmatch(r'total', low) and section == 'equity' and raw_weight is not None:
                value, units = weight_value(row[weight_col]); unit_types.add(units)
                equity_totals.append(value)
            continue
        kind = section_type(security) if not isin else None
        if kind and re.search(r'\btotal\b', low) and raw_weight is not None:
            value, units = weight_value(row[weight_col]); unit_types.add(units)
            if kind == 'equity':
                equity_totals.append(value)
            continue
        if kind:
            section, parent = section_transition(kind, security, parent)
            seen_equity = seen_equity or parent == 'equity'
            continue
        if not isin:
            if section == 'excluded':
                continue
            # Unknown weighted named headings could conceal a different asset
            # class. Totals and the recognised non-investment labels are safe.
            if security and raw_weight is not None and not re.search(r'total|net assets|nav|subtotal|sub total', low):
                try:
                    value, _ = weight_value(row[weight_col])
                except ValueError:
                    continue
                if value:
                    raise ValueError('Weighted row has no ISIN or recognised asset classification: ' + security)
            elif security and raw_weight is None and not re.search(r'total|net assets|nav|subtotal|sub total', low):
                section = None
            continue
        if section is None:
            raise ValueError('Security encountered before an explicit asset classification')
        if section != 'equity':
            continue
        if not valid_isin(isin):
            raise ValueError('Invalid physical-equity ISIN check digit: ' + isin)
        if re.search(r'\b(?:futures?|options?|swap|etf|fund|reits?|invits?)\b', low):
            raise ValueError('Non-physical-equity instrument within an equity section')
        weight, units = weight_value(row[weight_col]); unit_types.add(units)
        if weight < 0:
            raise ValueError('Negative physical-equity NAV weight')
        weighted_equity = True
        if not weight:
            omitted_zero += 1
            continue
        if not security or len(security) > 200:
            raise ValueError('Missing or oversized security name')
        sector = text(row[sector_col].value) if sector_col is not None and sector_col < len(row) else 'Unknown'
        sector = sector or 'Unknown'
        if isin in included:
            if included[isin]['sector'] != sector:
                raise ValueError('Conflicting sector labels for a duplicate ISIN')
            included[isin]['weight'] = round(included[isin]['weight'] + weight, 6)
        else:
            included[isin] = dict(isin=isin, name=security, sector=sector, weight=weight)
    if len(unit_types) != 1:
        raise ValueError('Mixed or missing percentage units')
    if not grand_totals or any(abs(value - 100) > .5 for value in grand_totals):
        raise ValueError('Portfolio grand total must reconcile to 100% NAV')
    if not seen_equity or not weighted_equity:
        return None, dict(**identity, portfolioDate=date.isoformat(), message='No physical listed-equity allocation in this portfolio.')
    total = round(sum(h['weight'] for h in included.values()), 6)
    if not included or len(included) > 2000 or not 0 < total <= 100.5:
        raise ValueError('Invalid physical-equity coverage')
    if equity_totals and min(abs(max(equity_totals) - total), abs(sum(equity_totals) - total)) > max(.5, len(included) * .011):
        raise ValueError('Physical-equity subtotal does not reconcile with included rows')
    snapshot = dict(schemeId=identity['id'], name=identity['name'], amc=source['amc'],
                    category=source.get('category', 'Physical equity portfolio'),
                    portfolioDate=date.isoformat(), includedNavWeight=total, belowPrecisionCount=omitted_zero,
                    policy=POLICY, source=dict(kind='official', publisher=source['amc'], url=document['url'],
                    indexUrl=source.get('indexUrl', source.get('pageUrl', document['url'])),
                    fetchedAt=fetched_at, sha256=source_sha, adapterVersion=1, worksheet=name),
                    holdings=sorted(included.values(), key=lambda h: (-h['weight'], h['isin'])))
    return (snapshot, identity['schemeCodes']), None


def workbook_sheets(raw, file_format):
    # Several AMC publishers label an OOXML workbook .xls. Verify the bytes,
    # rather than trusting the URL suffix or allowing arbitrary HTML tables.
    if raw.startswith(b'PK\x03\x04'):
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if len(members) > 1000 or sum(m.file_size for m in members) > MAX_EXPANDED:
                raise ValueError('Expanded workbook exceeds portfolio limits')
            if 'xl/workbook.xml' not in archive.namelist():
                raise ValueError('ZIP document is not an OOXML workbook')
        file_format = 'xlsx'
    elif not raw.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        raise ValueError('Document is not an XLS/XLSX workbook')
    if file_format == 'xlsx':
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            for sheet in workbook:
                rows, empty_run = [], 0
                for row in sheet.iter_rows(max_row=min(sheet.max_row, 25000), max_col=min(sheet.max_column, 120)):
                    converted = [Cell(c.value, c.number_format or 'General') for c in row]
                    empty_run = empty_run + 1 if not any(c.value is not None for c in converted) else 0
                    if empty_run > 100:
                        break
                    rows.append(converted)
                yield sheet.title, rows
        finally:
            workbook.close()
    elif file_format == 'xls':
        import xlrd  # Offline/development-only requirement; never a runtime server dependency.
        workbook = xlrd.open_workbook(file_contents=raw, formatting_info=True, logfile=io.StringIO())
        try:
            for sheet in workbook.sheets():
                rows, empty_run = [], 0
                for r in range(min(sheet.nrows, 25000)):
                    row = []
                    for c in range(min(sheet.ncols, 120)):
                        cell = sheet.cell(r, c)
                        value = xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode) if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                        xf = workbook.xf_list[sheet.cell_xf_index(r, c)]
                        fmt = workbook.format_map.get(xf.format_key)
                        row.append(Cell(value, fmt.format_str if fmt else 'General'))
                    empty_run = empty_run + 1 if not any(text(c.value) for c in row) else 0
                    if empty_run > 100:
                        break
                    rows.append(row)
                yield sheet.name, rows
        finally:
            workbook.release_resources()
    else:
        raise ValueError('Unsupported workbook format')


def read_workbooks(raw, file_format):
    if len(raw) > MAX_DOCUMENT:
        raise ValueError('Official document exceeds byte limit')
    if file_format != 'zip':
        yield '', raw, file_format
        return
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        members = [item for item in archive.infolist() if not item.is_dir()]
        if len(members) > MAX_ZIP_MEMBERS or sum(item.file_size for item in members) > MAX_EXPANDED:
            raise ValueError('Official ZIP exceeds expanded portfolio limits')
        for member in members:
            if member.flag_bits & 1 or member.file_size > MAX_DOCUMENT:
                raise ValueError('Encrypted or oversized ZIP member')
            path = member.filename.replace('\\', '/')
            if path.startswith('/') or '..' in path.split('/') or ':' in path:
                raise ValueError('Unsafe ZIP member path')
            kind = path.rsplit('.', 1)[-1].lower()
            if kind in ('xlsx', 'xls'):
                yield path, archive.read(member), kind


def adapt_uti_rows(rows, source, families):
    """UTI publishes numeric-text % NAV and an aggregate market-value total.

    Reconcile every original investment/cash market value to the explicitly
    named scheme total and every numeric row percentage to that denominator.
    Only then expose an equivalent 100% NAV total to the common parser. The
    source hash always remains the digest of the original downloaded bytes.
    """
    header, isin_col, weight_col, name_col, sector_col = find_columns(rows)
    identity = scheme_identity(rows[:header], source, families)
    markets = [i for i, c in enumerate(rows[header]) if re.search(r'market.*value', text(c.value), re.I)]
    if len(markets) != 1:
        raise ValueError('UTI disclosure needs one labelled market-value column')
    market_col = markets[0]
    def number(value):
        if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value):
            return float(value)
        if isinstance(value, str) and re.fullmatch(r'-?\d+(?:\.\d+)?', value.strip()):
            return float(value)
        raise ValueError('UTI market values and percentages require explicit numeric units')
    totals = []
    for index, row in enumerate(rows[header + 1:], header + 1):
        if len(row) <= max(name_col, market_col, weight_col):
            continue
        match = re.fullmatch(r'total\s*:\s*(.+)', text(row[name_col].value), re.I)
        if match and family_name(match[1]).replace(' ', '') == identity['family'].replace(' ', ''):
            totals.append((index, number(row[market_col].value)))
    if len(totals) != 1 or not totals[0][1] > 0:
        raise ValueError('UTI scheme aggregate market value is missing or ambiguous')
    end, nav = totals[0]
    normalized = [list(row) for row in rows[:end + 1]]
    market_sum, positions, section, parent = 0, 0, None, None
    for index, row in enumerate(normalized[header + 1:end], header + 1):
        if len(row) <= max(name_col, market_col, weight_col):
            continue
        label = text(row[name_col].value)
        if re.search(r'\btotal\b', label, re.I) or not label:
            continue
        kind = section_type(label) if not text(row[isin_col].value) else None
        if kind:
            section, parent = section_transition(kind, label, parent)
        market_raw = row[market_col].value
        if market_raw is None or text(market_raw).lower() in ('', 'nil', '-'):
            continue
        market = number(market_raw)
        market_sum += market; positions += 1
        weight_raw = row[weight_col].value
        if market == 0 and section == 'excluded':
            # Unlisted written-down positions may carry '*'; they remain
            # excluded and add exactly zero to the independently checked NAV.
            continue
        if weight_raw is None or text(weight_raw).lower() in ('', 'nil', '-'):
            if market and text(row[isin_col].value):
                raise ValueError('UTI priced security has no explicit NAV percentage')
            continue
        weight = number(weight_raw)
        if abs(weight - market / nav * 100) > .025:
            raise ValueError('UTI row NAV percentage does not reconcile with market value')
        row[weight_col] = Cell(weight, 'General')
    if not positions or abs(market_sum - nav) > max(.05, positions * .011):
        raise ValueError('UTI market values do not reconcile with the scheme NAV total')
    normalized[end][name_col] = Cell('Grand Total')
    normalized[end][weight_col] = Cell(100.0, 'General')
    return normalized


def parse_document(raw, source, document, catalogue, fetched_at, families=None):
    """Return verified snapshots plus explicit unsupported/rejected worksheets.

    The source digest refers to the downloaded workbook/ZIP bytes; immutable
    snapshot digests are computed by the publisher after deterministic JSON.
    """
    families = families if families is not None else catalogue_families(catalogue)
    digest = hashlib.sha256(raw).hexdigest()
    result = dict(funds=[], unsupported=[], errors=[])
    seen = set()
    for member, workbook_raw, kind in read_workbooks(raw, document['format']):
        for sheet_name, rows in workbook_sheets(workbook_raw, kind):
            # Non-portfolio cover/index worksheets are metadata, not failures.
            if not any(re.fullmatch(r'isin(?:\s+code)?', text(c.value), re.I) for row in rows[:35] for c in row):
                continue
            label = '/'.join(v for v in (member, sheet_name) if v)
            try:
                if source['id'] == 'uti':
                    rows = adapt_uti_rows(rows, source, families)
                fund, unsupported = parse_sheet(label, rows, source, document, families, fetched_at, digest)
                if fund and source['id'] == 'uti':
                    fund[0]['source']['adapter'] = 'uti-market-value-reconciliation-v1'
                    fund[0]['source']['reportStatus'] = 'Provisional and unaudited, as labelled by UTI.'
                if fund and document.get('disclosedUrl'):
                    fund[0]['source']['disclosedUrl'] = document['disclosedUrl']
                if fund and source.get('automaticRenewal') is False:
                    fund[0]['source']['automaticRenewal'] = False
                    fund[0]['source']['renewalStatus'] = source.get('renewalStatus', 'Future monthly selection requires manual verification.')
                identity = fund[0]['schemeId'] if fund else unsupported['id']
                if identity in seen:
                    raise ValueError('Duplicate scheme worksheets require an explicit adapter')
                seen.add(identity)
                if fund:
                    result['funds'].append(dict(snapshot=fund[0], schemeCodes=fund[1]))
                else:
                    result['unsupported'].append(unsupported)
            except ValueError as error:
                result['errors'].append(dict(worksheet=label, message=str(error)))
    if not result['funds'] and not result['unsupported'] and not result['errors']:
        raise ValueError('No recognisable portfolio worksheets in official document')
    return result
