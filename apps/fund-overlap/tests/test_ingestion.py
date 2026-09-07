import datetime as dt
import importlib.util
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('refresh', ROOT / 'ingestion' / 'refresh.py')
refresh = importlib.util.module_from_spec(spec)
spec.loader.exec_module(refresh)
FIXTURES = Path(__file__).parent / 'fixtures'
DATE = dt.date(2026, 7, 31)
URL = 'https://amc.ppfas.com/downloads/portfolio-disclosure/2026/PPFCF_PPFAS_Monthly_Portfolio_Report_July_31_2026.xlsx'

class IngestionTests(unittest.TestCase):
    def test_real_sources_and_golden_totals(self):
        expected = {'PPFCF': (62, 83.29), 'PPTSF': (41, 91.42), 'PPLCF': (55, 84.55), 'PPCHF': (11, 10.65), 'PPDAAF': (30, 26.59)}
        for code, (count, total) in expected.items():
            snapshot = refresh.parse_ppfas((FIXTURES / (code + '.xlsx')).read_bytes(), code, DATE, URL, '2026-09-07T00:00:00Z')
            self.assertEqual(len(snapshot['holdings']), count)
            self.assertEqual(snapshot['includedNavWeight'], total)
            self.assertEqual(len({h['isin'] for h in snapshot['holdings']}), count)
            self.assertTrue(all(refresh.valid_isin(h['isin']) for h in snapshot['holdings']))
        self.assertFalse(refresh.valid_isin('INE040A01035'))

    def test_rounding_and_duplicate_aggregation(self):
        raw = (FIXTURES / 'PPFCF.xlsx').read_bytes()
        result = refresh.parse_ppfas(raw, 'PPFCF', DATE, URL, 'now')
        self.assertGreater(result['belowPrecisionCount'], 0)
        self.assertEqual(next(h for h in result['holdings'] if h['isin'] == 'INE040A01034')['weight'], 7.55)
        workbook = load_workbook(io.BytesIO(raw))
        sheet = workbook['PPFCF']
        sheet['C8'] = sheet['C7'].value
        sheet['D8'] = sheet['D7'].value
        out = io.BytesIO(); workbook.save(out)
        merged = refresh.parse_ppfas(out.getvalue(), 'PPFCF', DATE, URL, 'now')
        self.assertEqual(next(h for h in merged['holdings'] if h['isin'] == 'INE040A01034')['weight'], 13.53)

    def test_date_and_bad_weight_rejection(self):
        raw = (FIXTURES / 'PPFCF.xlsx').read_bytes()
        with self.assertRaises(ValueError):
            refresh.parse_ppfas(raw, 'PPFCF', dt.date(2026, 8, 31), URL, 'now')
        workbook = load_workbook(io.BytesIO(raw)); sheet = workbook['PPFCF']; sheet['G7'] = -0.01
        out = io.BytesIO(); workbook.save(out)
        with self.assertRaises(ValueError):
            refresh.parse_ppfas(out.getvalue(), 'PPFCF', DATE, URL, 'now')
        sheet['G7'] = .0755
        sheet['C7'] = 'BROKEN-ISIN'
        out = io.BytesIO(); workbook.save(out)
        with self.assertRaises(ValueError):
            refresh.parse_ppfas(out.getvalue(), 'PPFCF', DATE, URL, 'now')

    def test_atomic_publication_noop_and_source_failure(self):
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp)
            refresh.refresh(output, FIXTURES, FIXTURES / 'index.html')
            before = (output / 'fund-index.json').read_bytes()
            self.assertEqual(len(json.loads(before)['funds']), 5)
            refresh.refresh(output, FIXTURES, FIXTURES / 'index.html')
            self.assertEqual((output / 'fund-index.json').read_bytes(), before)
            with patch.object(refresh, 'parse_ppfas', side_effect=ValueError('invalid source')):
                with self.assertRaises(ValueError):
                    refresh.refresh(output, FIXTURES, FIXTURES / 'index.html')
            self.assertEqual((output / 'fund-index.json').read_bytes(), before)

    def test_source_discovery_fails_closed(self):
        sources = refresh.discover((FIXTURES / 'index.html').read_text())
        self.assertEqual(set(sources), set(refresh.FUNDS))
        self.assertTrue(all(date == DATE for date, url in sources.values()))
        with self.assertRaises(ValueError):
            refresh.discover('<a href="https://example.com/unknown.xlsx">not a source</a>')

if __name__ == '__main__':
    unittest.main()
