import datetime as dt
import base64
import importlib.util
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'ingestion'))
import official_parser as parser
import official_refresh as refresh

DATE = dt.date(2026, 9, 30)
NOW = '2026-10-12T00:00:00+00:00'
SOURCE = dict(id='sample', amc='Sample Mutual Fund', indexUrl='https://sample.example/portfolios', enabled=True)
DOCUMENT = dict(url='https://sample.example/monthly-portfolio-2026-09.xlsx', portfolioDate='2026-09-30', format='xlsx', fileName='sample.xlsx')
CATALOGUE = [dict(schemeCode=101, schemeName='Sample Flexi Cap Fund - Direct Plan - Growth'),
             dict(schemeCode=102, schemeName='Sample Flexi Cap Fund - Regular Plan - IDCW'),
             dict(schemeCode=103, schemeName='Sample Growth Opportunities Fund - Direct Plan - Growth'),
             dict(schemeCode=104, schemeName='Sample Liquid Fund - Direct Plan - Growth'),
             dict(schemeCode=105, schemeName='Other Balanced Fund - Direct Plan - Growth')]


def workbook(name='Sample Flexi Cap Fund', units='excel-percent', section='Equity & Equity related', extra_sheet=False):
    book = Workbook(); sheet = book.active; sheet.title = 'PORTFOLIO'
    sheet.append([name]); sheet.append(['Monthly Portfolio as on September 30, 2026'])
    sheet.append(['Name of Instrument', 'ISIN', 'Industry / Rating', '% to Net Assets'])
    sheet.append([section])
    rows = [('HDFC Bank Limited', 'INE040A01034', 'Banks', 50), ('ITC Limited', 'INE154A01025', 'FMCG', 30)]
    for security, isin, sector, weight in rows:
        sheet.append([security, isin, sector, weight / 100 if units == 'excel-percent' else weight])
        if units == 'excel-percent': sheet.cell(sheet.max_row, 4).number_format = '0.00%'
    sheet.append(['Total Equity', None, None, .8 if units == 'excel-percent' else 80])
    sheet.append(['Debt Instruments'])
    sheet.append(['HDFC Debt Security', 'INE040A16HR5', 'AAA', .15 if units == 'excel-percent' else 15])
    sheet.append(['Cash and Cash Equivalents', None, None, .05 if units == 'excel-percent' else 5])
    sheet.append(['Grand Total', None, None, 1 if units == 'excel-percent' else 100])
    if units == 'excel-percent':
        for row in (7, 9, 10, 11): sheet.cell(row, 4).number_format = '0.00%'
    if extra_sheet:
        other = book.copy_worksheet(sheet); other.title = 'GROWTH'; other['A1'] = 'Sample Growth Opportunities Fund'
    out = io.BytesIO(); book.save(out); return out.getvalue()


def mutate(raw, coordinate, value, number_format=None):
    book = load_workbook(io.BytesIO(raw)); book.active[coordinate] = value
    if number_format: book.active[coordinate].number_format = number_format
    out = io.BytesIO(); book.save(out); return out.getvalue()


class OfficialParserTests(unittest.TestCase):
    def parse(self, raw, document=DOCUMENT):
        return parser.parse_document(raw, SOURCE, document, CATALOGUE, NOW)

    def test_exact_plan_families_and_no_interior_word_conflation(self):
        self.assertEqual(parser.family_name(CATALOGUE[0]['schemeName']), parser.family_name(CATALOGUE[1]['schemeName']))
        self.assertEqual(parser.family_name(CATALOGUE[2]['schemeName']), 'sample growth opportunities fund')
        self.assertNotEqual(parser.family_name('Sample Growth Fund - Direct Growth'), parser.family_name('Sample Fund - Direct Growth'))
        self.assertNotEqual(parser.family_name('Sample Small Cap Fund'), parser.family_name('Sample Mid Cap Fund'))
        self.assertEqual(parser.family_name('Sample - Growth Opportunities Fund - Direct Plan - Growth'), 'sample growth opportunities fund')
        result = self.parse(workbook())
        self.assertFalse(result['errors']); self.assertEqual(result['funds'][0]['schemeCodes'], [101, 102])

    def test_amc_ownership_requires_whole_brand_tokens(self):
        quant = dict(SOURCE, id='quant', amc='Quant Mutual Fund')
        quantum = dict(SOURCE, id='quantum', amc='Quantum Mutual Fund')
        self.assertTrue(parser.owns_family(quant, 'quant flexi cap fund'))
        self.assertFalse(parser.owns_family(quant, 'quantum long term equity value fund'))
        self.assertTrue(parser.owns_family(quantum, 'quantum long term equity value fund'))
        angel = dict(SOURCE, amc='Angel One Mutual Fund', namePrefixes=['Angel One', 'AngelOne'])
        self.assertTrue(parser.owns_family(angel, 'angelone nifty 50 etf'))

    def test_consolidated_per_sheet_holdings_exclude_debt_and_cash(self):
        result = self.parse(workbook(extra_sheet=True))
        self.assertFalse(result['errors']); self.assertEqual(len(result['funds']), 2)
        for fund in result['funds']:
            snapshot = fund['snapshot']
            self.assertEqual(snapshot['includedNavWeight'], 80)
            self.assertEqual({h['isin'] for h in snapshot['holdings']}, {'INE040A01034', 'INE154A01025'})
            self.assertEqual(snapshot['source']['worksheet'] in ('PORTFOLIO', 'GROWTH'), True)
            refresh.validate_snapshot(snapshot)
        # Explicit column units also support numeric NAV percentages.
        self.assertEqual(self.parse(workbook(units='numeric'))['funds'][0]['snapshot']['includedNavWeight'], 80)

    def test_identity_month_units_and_isin_fail_closed(self):
        cases = [('A1', 'Different Flexi Cap Fund', None),
                 ('A2', 'Monthly Portfolio as on August 31, 2026', None),
                 ('A4', 'Unknown strategic investments', None),
                 ('B5', 'INE040A01035', None),
                 ('D5', -.5, None), ('D5', 50, 'General'),
                 ('D11', .95, None), ('A5', 'Sample Gold ETF', None)]
        for coordinate, value, fmt in cases:
            with self.subTest(coordinate=coordinate, value=value):
                result = self.parse(mutate(workbook(), coordinate, value, fmt))
                self.assertFalse(result['funds']); self.assertTrue(result['errors'])

    def test_missing_percentage_header_and_subtotal_mismatch(self):
        for raw in (mutate(workbook(), 'D3', 'Market Value'), mutate(workbook(), 'D7', .65)):
            result = self.parse(raw); self.assertFalse(result['funds']); self.assertTrue(result['errors'])

    def test_debt_gold_and_fund_of_funds_are_explicitly_unsupported(self):
        for name in ('Sample Liquid Fund', 'Sample Gold Fund', 'Sample International Fund of Funds'):
            result = self.parse(workbook(name))
            self.assertFalse(result['funds']); self.assertFalse(result['errors']); self.assertEqual(len(result['unsupported']), 1)
            self.assertTrue(result['unsupported'][0]['message'])

    def test_equity_savings_and_regular_savings_hybrid_physical_equity_accepted(self):
        for name in ('Sample Equity Savings Fund', 'Sample Regular Savings Fund', 'Sample Monthly Income Fund'):
            result = self.parse(workbook(name))
            self.assertFalse(result['errors']); self.assertEqual(len(result['funds']), 1)
            self.assertEqual(result['funds'][0]['snapshot']['includedNavWeight'], 80)

    def test_unknown_section_cannot_inherit_equity_classification(self):
        raw = mutate(workbook(), 'A8', 'Alternative Investments')
        result = self.parse(raw)
        self.assertFalse(result['funds']); self.assertTrue(result['errors'])

    def test_listing_subsections_preserve_the_underlying_asset_class(self):
        book = load_workbook(io.BytesIO(workbook()))
        sheet = book.active
        sheet.insert_rows(5,2)
        sheet['A5'] = '(b) Unlisted'
        sheet['A6'] = '(a) Listed / awaiting listing on Stock Exchanges'
        out=io.BytesIO(); book.save(out)
        result = self.parse(out.getvalue())
        self.assertFalse(result['errors']); self.assertEqual(result['funds'][0]['snapshot']['includedNavWeight'], 80)
        # A genuine new debt asset section still changes the parent class.
        raw = mutate(out.getvalue(),'A6','Debt Instruments')
        result = self.parse(raw)
        self.assertFalse(result['funds'])

    def test_actual_amc_heading_column_aum_and_side_tables(self):
        book = load_workbook(io.BytesIO(workbook()))
        sheet = book.active
        sheet['A4'] = None; sheet['B4'] = 'Equity & Equity related'
        sheet['D3'] = '% to AUM'
        sheet['G4'] = '% to NAV'; sheet['G5'] = 50  # adjacent analysis table
        out = io.BytesIO(); book.save(out)
        result = self.parse(out.getvalue())
        self.assertFalse(result['errors']); self.assertEqual(result['funds'][0]['snapshot']['includedNavWeight'], 80)
        result=self.parse(mutate(workbook(),'A3','Company/Issuer/Instrument Name'))
        self.assertFalse(result['errors']); self.assertEqual(result['funds'][0]['snapshot']['includedNavWeight'],80)

    def test_month_heading_and_comma_spacing_require_exact_expected_month(self):
        for date_heading in ('Monthly Portfolio as on September 30,2026',
                             'Monthly Portfolio Statement for the period ended 30.September.2026',
                             'Monthly Portfolio Statement of Sample Flexi Cap Fund for September 2026'):
            result = self.parse(mutate(workbook(), 'A2', date_heading))
            self.assertFalse(result['errors']); self.assertEqual(len(result['funds']), 1)

    def test_xlsx_content_detection_and_inner_zip_limits(self):
        result = self.parse(workbook(), dict(DOCUMENT, format='xls'))
        self.assertFalse(result['errors']); self.assertEqual(len(result['funds']), 1)
        with patch.object(parser, 'MAX_EXPANDED', 100):
            with self.assertRaises(ValueError): self.parse(workbook())

    def test_identifier_length_is_bounded_for_long_source_and_scheme_names(self):
        source = dict(SOURCE, id='s' * 80)
        result = parser.parse_document(workbook('Sample ' + 'Long ' * 20 + 'Fund'), source, DOCUMENT, CATALOGUE, NOW)
        self.assertFalse(result['errors']); self.assertLessEqual(len(result['funds'][0]['snapshot']['schemeId']), 80)

    def test_equity_etf_heading_without_fund_is_supported(self):
        result = self.parse(workbook('Sample Nifty 50 ETF'))
        self.assertFalse(result['errors']); self.assertEqual(len(result['funds']), 1)
        self.assertEqual(result['funds'][0]['snapshot']['includedNavWeight'], 80)

    def test_dated_manual_discovery_does_not_claim_automatic_renewal(self):
        source=dict(SOURCE,automaticRenewal=False,renewalStatus='Next month needs manual source verification.')
        result=parser.parse_document(workbook(),source,DOCUMENT,CATALOGUE,NOW)
        self.assertFalse(result['funds'][0]['snapshot']['source']['automaticRenewal'])
        self.assertIn('manual',result['funds'][0]['snapshot']['source']['renewalStatus'])

    def test_uti_market_value_adapter_requires_full_nav_reconciliation(self):
        source = dict(SOURCE, id='uti', amc='UTI Mutual Fund')
        book = load_workbook(io.BytesIO(workbook('UTI Flexi Cap Fund', units='numeric')))
        sheet = book.active; sheet['E3'] = 'Market-Value'
        # Physical equity 500+300, debt150 and cash50 sum to disclosed NAV1000.
        for row, value in ((5,500),(6,300),(9,150),(10,50)): sheet.cell(row,5,value)
        sheet['A11'] = 'TOTAL : UTI Flexi Cap Fund'; sheet['D11'] = None; sheet['E11'] = 1000
        sheet['D5'] = '50.00'; sheet['D6'] = '30.00'
        out = io.BytesIO(); book.save(out)
        result = parser.parse_document(out.getvalue(), source, DOCUMENT, CATALOGUE, NOW)
        self.assertFalse(result['errors']); self.assertEqual(result['funds'][0]['snapshot']['includedNavWeight'], 80)
        self.assertIn('uti-market-value', result['funds'][0]['snapshot']['source']['adapter'])
        for coordinate, value in (('E11',900),('D5','49.00'),('E6',310)):
            result = parser.parse_document(mutate(out.getvalue(),coordinate,value),source,DOCUMENT,CATALOGUE,NOW)
            self.assertFalse(result['funds']); self.assertTrue(result['errors'])

    def test_zip_with_multiple_workbooks_and_traversal_rejected(self):
        raw = io.BytesIO()
        with zipfile.ZipFile(raw, 'w') as archive:
            archive.writestr('flexi.xlsx', workbook())
            archive.writestr('growth.xlsx', workbook('Sample Growth Opportunities Fund'))
        document = dict(DOCUMENT, format='zip', url='https://sample.example/portfolio.zip')
        result = self.parse(raw.getvalue(), document)
        self.assertFalse(result['errors']); self.assertEqual(len(result['funds']), 2)
        with patch.object(parser, 'MAX_ZIP_MEMBERS', 1):
            with self.assertRaises(ValueError): self.parse(raw.getvalue(), document)
        unsafe = io.BytesIO()
        with zipfile.ZipFile(unsafe, 'w') as archive: archive.writestr('../escape.xlsx', workbook())
        with self.assertRaises(ValueError): self.parse(unsafe.getvalue(), document)

    def test_legacy_xls_fixture_has_same_verified_contract(self):
        fixture = Path(__file__).parent / 'fixtures' / 'official-generic.xls.b64'
        result = self.parse(base64.b64decode(fixture.read_text()), dict(DOCUMENT, format='xls', url='https://sample.example/portfolio.xls'))
        self.assertFalse(result['errors']); self.assertEqual(result['funds'][0]['snapshot']['includedNavWeight'], 80)

    def test_legacy_xls_style_tail_is_bounded_without_rejecting_the_portfolio(self):
        fixture = Path(__file__).parent / 'fixtures' / 'official-overformatted.xls.b64'
        result = self.parse(base64.b64decode(fixture.read_text()), dict(DOCUMENT, format='xls', url='https://sample.example/portfolio.xls'))
        self.assertFalse(result['errors']); self.assertEqual(result['funds'][0]['snapshot']['includedNavWeight'], 80)

    def test_real_amc_golden_excerpts_and_provenance(self):
        fixtures = Path(__file__).parent / 'fixtures'
        for amc in ('abakkus', 'hdfc', 'uti'):
            metadata = json.loads((fixtures / ('official-' + amc + '-excerpt.json')).read_text())
            raw = base64.b64decode((fixtures / ('official-' + amc + '-excerpt.xlsx.b64')).read_text())
            result = parser.parse_document(raw, metadata['source'], metadata['document'], metadata['catalogue'], NOW)
            self.assertFalse(result['errors']); self.assertEqual(len(result['funds']), 1)
            snapshot = result['funds'][0]['snapshot']
            self.assertEqual(len(snapshot['holdings']), metadata['expected']['count'])
            self.assertEqual(snapshot['includedNavWeight'], metadata['expected']['includedNavWeight'])
            self.assertEqual(set(result['funds'][0]['schemeCodes']), {v['schemeCode'] for v in metadata['catalogue']})
            self.assertRegex(metadata['derivedFromSha256'], '^[a-f0-9]{64}$')
            refresh.validate_snapshot(snapshot)


class OfficialRefreshTests(unittest.TestCase):
    def run_refresh(self, output, sources, fetch):
        return refresh.refresh(output, sources, CATALOGUE, as_of=DATE + dt.timedelta(days=12),
                               fetch_bytes=fetch, discover=lambda source, fetch, date: [DOCUMENT])

    def test_partial_failure_preserves_last_good_and_all_national_statuses(self):
        other = dict(SOURCE, id='other', amc='Other Mutual Fund')
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)
            first = self.run_refresh(output, [SOURCE, other], lambda url: workbook())
            self.assertEqual(len(first['funds']), 1)
            entry = first['funds'][0]; snapshot = refresh.snapshot_path(output, entry).read_bytes()
            self.assertEqual(len(first['coverageReport']['schemes']), len(CATALOGUE))
            self.assertEqual(first['coverageReport']['summary']['unsupportedSchemes'], 1)
            def bad_fetch(url): raise OSError('Provider unavailable')
            second = self.run_refresh(output, [SOURCE, other], bad_fetch)
            self.assertEqual(second['funds'][0]['holdingsPath'], entry['holdingsPath'])
            self.assertEqual(refresh.snapshot_path(output, second['funds'][0]).read_bytes(), snapshot)
            self.assertEqual(second['funds'][0]['coverageStatus'], 'stale')
            self.assertEqual(second['coverageReport']['summary']['staleSchemes'], 2)
            self.assertEqual(json.loads((output / 'coverage.json').read_text())['generation'], second['generation'])
            self.assertEqual(json.loads((output / 'official-index.json').read_text())['coverageReport'], second['coverageReport'])
            third = self.run_refresh(output, [SOURCE], lambda url: workbook())
            self.assertNotIn('coverageStatus', third['funds'][0])

    def test_clean_source_omission_retains_missing_fund_as_stale(self):
        with tempfile.TemporaryDirectory() as temporary:
            output=Path(temporary)
            first=self.run_refresh(output,[SOURCE],lambda url:workbook(extra_sheet=True))
            self.assertEqual(len(first['funds']),2)
            missing=next(e for e in first['funds'] if 'growth-opportunities' in e['id'])
            second=self.run_refresh(output,[SOURCE],lambda url:workbook())
            funds={e['id']:e for e in second['funds']}
            self.assertEqual(funds[missing['id']]['holdingsPath'],missing['holdingsPath'])
            self.assertEqual(funds[missing['id']]['coverageStatus'],'stale')
            self.assertEqual(second['coverageReport']['sources'][0]['status'],'partial')
            self.assertEqual(second['coverageReport']['summary']['staleSchemes'],1)
            verified=next(r for r in second['coverageReport']['schemes'] if r['status']=='verified')
            self.assertNotIn('last valid snapshot retained',verified['message'].lower())

    def test_quant_and_quantum_have_unambiguous_coverage_owners(self):
        sources=[dict(SOURCE,id='quant',amc='Quant Mutual Fund',enabled=False),dict(SOURCE,id='quantum',amc='Quantum Mutual Fund',enabled=False)]
        catalogue=[dict(schemeCode=123456,schemeName='Quantum Long Term Equity Value Fund - Direct Plan - Growth')]
        with tempfile.TemporaryDirectory() as temporary:
            index=refresh.refresh(Path(temporary),sources,catalogue,as_of=DATE,fetch_bytes=lambda url:b'',discover=lambda *args:[])
            self.assertEqual(index['coverageReport']['schemes'][0]['amc'],'Quantum Mutual Fund')

    def test_successful_other_source_survives_one_failed_source(self):
        other = dict(SOURCE, id='other', amc='Other Mutual Fund')
        docs = dict(sample=DOCUMENT, other=dict(DOCUMENT, url='https://other.example/portfolio.xlsx'))
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)
            def discover(source, fetch, date): return [docs[source['id']]]
            def initial(url): return workbook('Other Balanced Fund') if 'other.example' in url else workbook()
            first = refresh.refresh(output, [SOURCE, other], CATALOGUE, as_of=DATE, fetch_bytes=initial, discover=discover)
            self.assertEqual(len(first['funds']), 2)
            def updated(url):
                if 'sample.example' in url: raise OSError('Sample failed')
                return mutate(workbook('Other Balanced Fund'), 'D5', .49)
            # Correct a source row and matching subtotal while retaining 100%.
            changed = mutate(mutate(workbook('Other Balanced Fund'), 'D5', .49), 'D7', .79)
            def updated(url):
                if 'sample.example' in url: raise OSError('Sample failed')
                return changed
            second = refresh.refresh(output, [SOURCE, other], CATALOGUE, as_of=DATE, fetch_bytes=updated, discover=discover)
            self.assertEqual(len(second['funds']), 2)
            funds = {f['sourceId']: f for f in second['funds']}
            self.assertEqual(funds['sample']['coverageStatus'], 'stale')
            self.assertEqual(funds['other']['includedNavWeight'], 79)

    def test_date_regression_and_bad_previous_digest_do_not_replace_good_snapshot(self):
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)
            first = self.run_refresh(output, [SOURCE], lambda url: workbook())
            entry = first['funds'][0]
            older = dict(DOCUMENT, portfolioDate='2026-08-31')
            old_raw = mutate(workbook(), 'A2', 'Monthly Portfolio as on August 31, 2026')
            second = refresh.refresh(output, [SOURCE], CATALOGUE, as_of=DATE, fetch_bytes=lambda url: old_raw,
                                     discover=lambda *args: [older])
            self.assertEqual(second['funds'][0]['holdingsPath'], entry['holdingsPath'])
            before = (output / 'official-index.json').read_bytes()
            refresh.snapshot_path(output, entry).write_bytes(b'{}')
            with self.assertRaises(ValueError): self.run_refresh(output, [SOURCE], lambda url: workbook())
            self.assertEqual((output / 'official-index.json').read_bytes(), before)

    def test_invalid_catalogue_and_unbounded_fixture_names_fail(self):
        with self.assertRaises(ValueError): refresh.validate_catalogue(CATALOGUE + [CATALOGUE[0]])
        with self.assertRaises(ValueError): refresh.fixture_document('.', SOURCE, dict(DOCUMENT, fileName='../outside.xlsx'))

    def test_snapshot_verification_rejects_invalid_business_values(self):
        snapshot = parser.parse_document(workbook(), SOURCE, DOCUMENT, CATALOGUE, NOW)['funds'][0]['snapshot']
        for modify in (lambda s: s['holdings'][0].update(isin='INE040A01035'),
                       lambda s: s['holdings'][0].update(weight=-1),
                       lambda s: s.update(includedNavWeight=100),
                       lambda s: s['source'].update(url='http://sample.example/portfolio.xlsx')):
            copy = json.loads(json.dumps(snapshot)); modify(copy)
            with self.assertRaises(ValueError): refresh.validate_snapshot(copy)

    def test_same_byte_source_retains_changed_provenance_metadata(self):
        with tempfile.TemporaryDirectory() as temporary:
            output=Path(temporary)
            first=self.run_refresh(output,[SOURCE],lambda url:workbook())
            original=refresh.parse_official
            def add_metadata(*args,**kwargs):
                result=original(*args,**kwargs)
                result['funds'][0]['snapshot']['source']['reportStatus']='Provisional and unaudited'
                return result
            with patch.object(refresh,'parse_official',side_effect=add_metadata):
                second=self.run_refresh(output,[SOURCE],lambda url:workbook())
            self.assertEqual(first['funds'][0]['sourceSha256'],second['funds'][0]['sourceSha256'])
            self.assertNotEqual(first['funds'][0]['holdingsPath'],second['funds'][0]['holdingsPath'])
            snapshot=json.loads(refresh.snapshot_path(output,second['funds'][0]).read_bytes())
            self.assertEqual(snapshot['source']['reportStatus'],'Provisional and unaudited')

    def test_index_is_last_atomic_commit_on_report_write_failure(self):
        with tempfile.TemporaryDirectory() as temporary:
            output = Path(temporary)
            self.run_refresh(output, [SOURCE], lambda url: workbook())
            before = (output / 'official-index.json').read_bytes()
            writer = refresh.atomic_write
            def fail_report(path, raw):
                if path.name == 'coverage.json': raise OSError('Disk write failed')
                writer(path, raw)
            with patch.object(refresh, 'atomic_write', side_effect=fail_report):
                with self.assertRaises(OSError): self.run_refresh(output, [SOURCE], lambda url: workbook())
            self.assertEqual((output / 'official-index.json').read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
